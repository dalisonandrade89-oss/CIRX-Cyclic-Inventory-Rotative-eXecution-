#!/usr/bin/env python3
"""Atualiza os dados embutidos (const EMBED) do webapp a partir da planilha de contagens.

Replica exatamente o carregador do botão "Atualizar contagens" (mesmas folhas, colunas,
linhas mínimas e datas como número de série do Excel), por isso o resultado é idêntico
ao que o navegador geraria ao carregar o mesmo ficheiro.

Uso:
    pip install openpyxl
    python3 atualizar_dados.py "Contagens de stock 2026.html" "Contagens 2026.xlsx" [AAAA-MM-DD]

A data (opcional, por omissão hoje) vai para APP.dataDate, mostrada no rodapé.
Depois de rodar: suba o PATCH da versão (APP.version) e registe no CHANGELOG.md.
"""
import sys, re, json, datetime
import openpyxl

EPOCH = datetime.datetime(1899, 12, 30)
SHEETS = re.compile(r'^(Contagens 2026 \(C\d+\)|Aud \d+|Aud oc \d+)$')


def cell(v):
    if isinstance(v, str):
        return None if v.strip() == '' else v
    if isinstance(v, datetime.datetime):
        d = v - EPOCH
        x = d.days + d.seconds / 86400
        return int(x) if x == int(x) else x
    if isinstance(v, datetime.date):
        return (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def build_embed(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        name = ws.title
        if not SHEETS.match(name):
            continue
        is_c = name.startswith('Contagens')
        is_oc = name.startswith('Aud oc')
        nc = 21 if is_c else 7 if is_oc else 6
        min_r = 2 if is_c else 3 if is_oc else 2
        rows = []
        for ln in range(ws.min_row, ws.max_row + 1):
            if ln < min_r:
                continue
            v = [cell(ws.cell(row=ln, column=c + 1).value) for c in range(nc)]
            if is_c:
                if v[0] is None and v[2] is None:
                    continue
            elif v[0] is None:
                continue
            rows.append([ln] + v)
        out[name] = rows
    if not any(k.startswith('Contagens') for k in out):
        sys.exit('Erro: não encontrei folhas "Contagens 2026 (C01)", "(C02)"…')
    return out


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    html_path, xlsx_path = sys.argv[1], sys.argv[2]
    data_date = sys.argv[3] if len(sys.argv) > 3 else datetime.date.today().isoformat()
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', data_date):
        sys.exit('Erro: a data deve estar no formato AAAA-MM-DD.')

    embed = build_embed(xlsx_path)
    payload = json.dumps(embed, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')

    html = open(html_path, encoding='utf-8').read()
    lines = html.split('\n')
    idx = [i for i, l in enumerate(lines) if l.startswith('const EMBED={')]
    if len(idx) != 1:
        sys.exit('Erro: não encontrei exatamente uma linha "const EMBED=" no HTML.')
    lines[idx[0]] = 'const EMBED=' + payload + ';'
    html = '\n'.join(lines)
    html, n = re.subn(r"dataDate:'\d{4}-\d{2}-\d{2}'", f"dataDate:'{data_date}'", html, count=1)
    if n != 1:
        sys.exit('Erro: não encontrei APP.dataDate no HTML.')
    open(html_path, 'w', encoding='utf-8').write(html)

    for k, v in embed.items():
        if k.startswith('Contagens'):
            print(f'{k}: {len(v)} linhas, {sum(1 for r in v if r[6] is not None)} com data')
    print(f'OK. APP.dataDate = {data_date}. Lembre-se de subir APP.version e registar no CHANGELOG.md.')


if __name__ == '__main__':
    main()
